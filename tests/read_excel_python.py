import openpyxl
import sys
import os

# Percorso del file Excel
excel_file = os.path.join(os.path.dirname(__file__), '..', 'template', 'Orario Allievi MusicAll.xlsx')

try:
    # Carica il workbook
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    
    # Lista dei fogli (giorni)
    sheets = wb.sheetnames
    print(f"Fogli disponibili: {sheets}")
    
    # Leggi il foglio MARTEDI
    if 'MARTEDI' in sheets:
        ws = wb['MARTEDI']
        print(f"\n=== MARTEDI ===\n")
        
        # Stampa prime 50 righe e 10 colonne per capire la struttura
        for row in range(1, 51):
            for col in range(1, 11):  # Colonne A-J
                cell = ws.cell(row=row, column=col)
                if cell.value:
                    col_letter = openpyxl.utils.get_column_letter(col)
                    print(f"{col_letter}{row}: {cell.value}")
    
    wb.close()
    
except Exception as e:
    print(f"Errore: {e}")
    sys.exit(1)