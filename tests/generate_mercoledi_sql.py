import openpyxl
import sys
import os

# Percorso del file Excel
excel_file = os.path.join(os.path.dirname(__file__), '..', 'template', 'Orario Soci MusicAll.xlsx')

# Mapping aule per MERCOLEDI (stesso del MARTEDI)
aule_map = {
    'C': ('MIDI', 'Chitarra'),
    'D': ('PIANO', 'Piano e Canto'), 
    'E': ('MAGNA', 'Canto'),
    'F': ('JAZZ', 'Batteria'),
    'G': ('POP', 'Canto'),
    'H': ('ROCK', 'Canto')
}

try:
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    ws = wb['MERCOLEDI']
    
    print("=== GENERAZIONE SQL MERCOLEDI ===\n")
    
    sql_statements = []
    sql_statements.append("-- =============================================")
    sql_statements.append("-- LEZIONI MERCOLEDI")
    sql_statements.append("-- Generato automaticamente dall'Excel")
    sql_statements.append("-- =============================================\n")
    
    lezioni_count = 0
    
    # Leggi tutte le righe con orari e nomi
    for row in range(10, 45):
        for col in range(3, 9):  # Colonne C-H
            cell = ws.cell(row=row, column=col)
            if not cell.value:
                continue
                
            value = str(cell.value).strip()
            
            # Salta celle non utili
            if any(skip in value.upper() for skip in ['DOCENTE', 'AULA', 'NOTE', 'PAUSA', 'NO ']):
                continue
            
            # Se contiene orario e nome allievo
            if '/' in value or ('-' in value and any(c.isalpha() for c in value)):
                col_letter = openpyxl.utils.get_column_letter(col)
                aula_nome, materia = aule_map.get(col_letter, (None, None))
                
                if not aula_nome:
                    continue
                
                # Estrai orario
                orario = None
                allievo = None
                
                parts = value.split()
                for part in parts:
                    if '/' in part or ('-' in part and ':' in part):
                        orario = part.replace('/', '-')
                        break
                
                if orario:
                    # Il resto è l'allievo
                    allievo_text = value.replace(orario, '').strip()
                    # Rimuovi parentesi e note
                    allievo_text = allievo_text.split('(')[0].strip()
                    
                    # Prendi solo cognome (prima parola maiuscola significativa)
                    words = allievo_text.split()
                    for word in words:
                        if word.isupper() and len(word) > 2:
                            allievo = word
                            break
                    
                    if allievo and 'PROVA' not in value.upper():
                        # Determina docente dalla riga 6-7
                        docente_cell = ws.cell(row=6, column=col).value or ws.cell(row=7, column=col).value
                        if docente_cell:
                            docente = str(docente_cell).split()[0].upper()
                        else:
                            docente = 'UNKNOWN'
                        
                        # Formatta orario
                        orario_parts = orario.replace('.', ':').split('-')
                        ora_inizio = orario_parts[0].strip()
                        ora_fine = orario_parts[1].strip() if len(orario_parts) > 1 else ora_inizio
                        
                        # Formato HH:MM
                        if len(ora_inizio.split(':')[0]) == 1:
                            ora_inizio = '0' + ora_inizio
                        if len(ora_fine.split(':')[0]) == 1:
                            ora_fine = '0' + ora_fine
                        
                        sql = f"""INSERT INTO lezioni (giorno_settimana, ora_inizio, ora_fine, id_aula, id_allievo, id_docente, id_materia)
SELECT 'MERCOLEDÌ', '{ora_inizio}', '{ora_fine}',
       (SELECT id FROM aule WHERE UPPER(nome) LIKE '%{aula_nome}%' LIMIT 1),
       (SELECT s.id FROM soci s JOIN persone p ON s.persona_id = p.id WHERE UPPER(p.cognome) LIKE '%{allievo}%' LIMIT 1),
       (SELECT id FROM docenti WHERE UPPER(cognome) LIKE '%{docente}%' LIMIT 1),
       (SELECT id FROM materie WHERE UPPER(nome) LIKE '%{materia.split()[0]}%' LIMIT 1);"""
                        
                        sql_statements.append(sql)
                        lezioni_count += 1
                        print(f"  + {aula_nome:8} {ora_inizio}-{ora_fine:5} {allievo:20} ({docente})")
    
    # Scrivi file
    output_file = os.path.join(os.path.dirname(__file__), '..', 'database', 'insert_lezioni_MERCOLEDI.sql')
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write('\n\n'.join(sql_statements))
    
    print(f"\nGenerati {lezioni_count} INSERT SQL per MERCOLEDI")
    print(f"File: {output_file}")
    
    wb.close()

except Exception as e:
    print(f"ERRORE: {e}")
    import traceback
    traceback.print_exc()
    sys.exit(1)